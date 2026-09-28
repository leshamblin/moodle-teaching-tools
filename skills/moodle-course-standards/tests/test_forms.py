from __future__ import annotations
import os
import sys
import tempfile
import warnings

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'scripts'))
import openpyxl
import forms
import rules

warnings.filterwarnings('ignore', message='Data Validation extension')


def _fill(course, results, term=''):
    d = tempfile.mkdtemp()
    out = os.path.join(d, 'x.xlsx')
    forms.fill_one(course, results, forms.TEMPLATE, out, term)
    return openpyxl.load_workbook(out)['Sheet1']


COURSE = {'id': 1, 'shortname': 'MKT 510 (631)',
          'fullname': 'MKT 510 (631) Fall 2026 Marketing', 'instructor': 'Smith'}


def test_fill_writes_verdict_into_column_b():
    ws = _fill(COURSE, [rules.Result(8, rules.NO, 'no banner'),
                        rules.Result(9, rules.YES, 'named sections')])
    assert ws.cell(8, 2).value == 'No'
    assert ws.cell(9, 2).value == 'Yes'


def test_evidence_goes_in_comment_column_and_recommend_text_survives():
    before = openpyxl.load_workbook(forms.TEMPLATE)['Sheet1'].cell(8, 3).value
    ws = _fill(COURSE, [rules.Result(8, rules.NO, 'no banner image')])
    assert ws.cell(8, 6).value == 'no banner image'
    assert ws.cell(8, 3).value == before


def test_fill_writes_header_block():
    ws = _fill(COURSE, [], term='Fall 2026')
    assert 'MKT 510' in ws.cell(1, 2).value
    assert ws.cell(2, 2).value == 'Fall 2026'
    assert ws.cell(3, 2).value == 'Smith'


def test_unknown_is_written_as_a_comment_not_a_dropdown_value():
    ws = _fill(COURSE, [rules.Result(43, rules.UNKNOWN, 'syllabus not readable: scan')])
    assert ws.cell(43, 2).value in (None, '')
    assert 'could not check' in ws.cell(43, 6).value.lower()
