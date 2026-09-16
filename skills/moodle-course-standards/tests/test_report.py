from __future__ import annotations
import os
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'scripts'))
import openpyxl
import report
import rules

COURSE_A = {'id': 1, 'shortname': 'MBA 520 (631)', 'fullname': 'MBA 520 (631) Spring 2026',
            'sections': ['631'], 'instructor': 'Smith'}
COURSE_B = {'id': 2, 'shortname': 'MBA 507 (632)', 'fullname': 'MBA 507 (632) Spring 2026',
            'sections': ['632'], 'instructor': 'Jones'}

RES_A = [rules.Result(8, rules.NO, 'no course image'),
         rules.Result(9, rules.NO, 'default section names'),
         rules.Result(10, rules.YES, 'consistent'),
         rules.Result(43, rules.UNKNOWN, 'syllabus not readable: scan')]
RES_B = [rules.Result(8, rules.YES, 'image set'),
         rules.Result(9, rules.YES, 'named'),
         rules.Result(10, rules.YES, 'consistent'),
         rules.Result(43, rules.NO, 'no accessibility statement')]


def test_summarise_counts_only_confirmed_no():
    s = report.summarise(COURSE_A, RES_A)
    assert s['no'] == 2
    assert s['unknown'] == 1
    assert s['yes'] == 1


def test_summarise_names_worst_failures():
    s = report.summarise(COURSE_A, RES_A)
    assert 'Course Banner' in s['top_failures']


def test_ranking_sorts_worst_first():
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, 'r.xlsx')
        report.write_workbook([{'course': COURSE_B, 'results': RES_B},
                               {'course': COURSE_A, 'results': RES_A}], out)
        wb = openpyxl.load_workbook(out)
        ws = wb['Ranking']
        assert ws.cell(2, 1).value == 'MBA 520 (631)'
        assert ws.cell(3, 1).value == 'MBA 507 (632)'


def test_unknown_does_not_push_a_course_down_the_ranking():
    """A course with 1 No and 3 Unknown must rank above one with 2 No."""
    res_unknown = [rules.Result(8, rules.NO, 'x')] + [
        rules.Result(r, rules.UNKNOWN, 'unreadable') for r in (41, 42, 43)]
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, 'r.xlsx')
        report.write_workbook([{'course': COURSE_B, 'results': res_unknown},
                               {'course': COURSE_A, 'results': RES_A}], out)
        ws = openpyxl.load_workbook(out)['Ranking']
        assert ws.cell(2, 1).value == 'MBA 520 (631)'


def test_three_sheets_exist():
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, 'r.xlsx')
        report.write_workbook([{'course': COURSE_A, 'results': RES_A}], out)
        wb = openpyxl.load_workbook(out)
        assert wb.sheetnames == ['Ranking', 'Detail', 'Evidence']


def test_detail_has_one_row_per_course_per_result():
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, 'r.xlsx')
        report.write_workbook([{'course': COURSE_A, 'results': RES_A},
                               {'course': COURSE_B, 'results': RES_B}], out)
        ws = openpyxl.load_workbook(out)['Detail']
        assert ws.max_row == 1 + len(RES_A) + len(RES_B)


def test_evidence_sheet_only_carries_no_and_unknown():
    with tempfile.TemporaryDirectory() as d:
        out = os.path.join(d, 'r.xlsx')
        report.write_workbook([{'course': COURSE_A, 'results': RES_A}], out)
        ws = openpyxl.load_workbook(out)['Evidence']
        assert ws.max_row == 1 + 3
