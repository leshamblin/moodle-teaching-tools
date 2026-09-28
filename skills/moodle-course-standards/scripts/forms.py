"""Fill one copy of the supplied checklist per course.

Writes into the original workbook so its Yes/No/Somewhat/N/A dropdowns and its
pre-written RECOMMEND text in column C survive. The verdict goes in column B
and the evidence in column F, the reviewer comment column. Off by default: the
stated deliverable is the tool, not the paperwork.
"""
from __future__ import annotations

import os
import re
from typing import Dict, List

import openpyxl

import rules

FORM_VALUES = (rules.YES, rules.NO, rules.SOMEWHAT, rules.NA)
VERDICT_COL = 2
COMMENT_COL = 6
TEMPLATE = os.path.join(os.path.dirname(__file__), '..', 'templates',
                        'Course-Review-Checklist.xlsx')


def _safe(name: str) -> str:
    return re.sub(r'[^A-Za-z0-9._-]+', '-', name).strip('-')


def fill_one(course: Dict, results: List, template: str, out_path: str,
             term: str = '') -> str:
    wb = openpyxl.load_workbook(template)
    ws = wb['Sheet1']

    ws.cell(1, 2).value = course.get('fullname', '')
    ws.cell(2, 2).value = term
    ws.cell(3, 2).value = course.get('instructor', '')
    ws.cell(4, 2).value = 'Automated audit, moodle-course-standards'

    for r in results:
        if r.verdict in FORM_VALUES:
            ws.cell(r.row, VERDICT_COL).value = r.verdict
            ws.cell(r.row, COMMENT_COL).value = r.evidence
        else:
            # Unknown is not one of the form's values. Leave the dropdown blank
            # so a human fills it, and say why in the comment.
            ws.cell(r.row, VERDICT_COL).value = None
            ws.cell(r.row, COMMENT_COL).value = 'Could not check automatically: %s' % r.evidence

    os.makedirs(os.path.dirname(out_path) or '.', exist_ok=True)
    wb.save(out_path)
    return out_path


def write_all(rows: List[Dict], out_dir: str, term: str = '') -> int:
    n = 0
    for row in rows:
        course = row['course']
        name = _safe(course.get('shortname') or str(course.get('id')))
        fill_one(course, row['results'], TEMPLATE,
                 os.path.join(out_dir, '%s.xlsx' % name), term)
        n += 1
    return n
